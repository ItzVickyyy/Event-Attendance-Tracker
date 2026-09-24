"""Student Import Service - XLSX parsing and staging"""

import re
from typing import Any

from openpyxl import load_workbook
from sqlmodel import Session

from app.models import ImportBatch, ImportValidationStatus, StudentImportRecord


class StudentImportService:
    """Service for parsing XLSX student imports and staging records"""

    def __init__(self, session: Session):
        self.session = session
        # Populated by parse_student_import() each time it runs; holds the
        # 3C-04 Summary Reconciliation result for the most recently parsed
        # workbook. Kept off the parse_student_import() return value so the
        # existing List[Dict[str, Any]] contract (and every caller/test that
        # depends on it) is left unchanged. See get_summary_reconciliation().
        self._last_summary_reconciliation: dict[str, Any] | None = None

    def parse_student_import(
        self, import_batch: ImportBatch, xlsx_file: bytes
    ) -> list[dict[str, Any]]:
        """Parse XLSX workbook and create StudentImportRecord staging rows.

        Returns:
            List of parsed row data with validation results
        """
        import io

        try:
            workbook = load_workbook(
                filename=io.BytesIO(xlsx_file), read_only=True, data_only=True
            )
        except Exception as e:
            raise ValueError(f"Failed to parse XLSX file: {str(e)}")

        parsed_rows = []
        processed_section_sheets: list[str] = []
        summary_sheet_rows: list[tuple[Any, ...]] | None = None

        for sheet_name in workbook.sheetnames:
            if sheet_name.strip().lower() == "summary":
                # The Summary sheet is a reconciliation/validation source,
                # not a source of student records (Phase-3 mapping, section
                # 10). Capture its rows for reconciliation below, but never
                # feed it into parsed_rows / staging.
                if summary_sheet_rows is None:
                    summary_sheet_rows = list(
                        workbook[sheet_name].iter_rows(values_only=True)
                    )
                continue

            if not self._validate_sheet_name(sheet_name):
                continue

            sheet = workbook[sheet_name]
            sheet_rows = list(sheet.iter_rows(values_only=True))

            if not sheet_rows:
                continue

            processed_section_sheets.append(sheet_name)

            header_index = self._find_header_row(sheet_rows)
            header_row = sheet_rows[header_index]
            data_rows = sheet_rows[header_index + 1 :]

            for row_idx, row in enumerate(data_rows, start=header_index + 2):
                if self._is_empty_row(row):
                    continue

                row_data = self._extract_row_data(
                    sheet_name, row_idx, row, header_row
                )
                parsed_rows.append(row_data)

        workbook.close()

        self._validate_and_detect_conflicts(parsed_rows)

        self.create_staging_records(import_batch, parsed_rows)

        self._last_summary_reconciliation = self._reconcile_summary_sheet(
            summary_sheet_rows, parsed_rows, processed_section_sheets
        )

        return parsed_rows

    def get_summary_reconciliation(self) -> dict[str, Any] | None:
        """Return the 3C-04 Summary Reconciliation result for the most
        recently parsed workbook, or None if parse_student_import() has not
        been run yet on this service instance.

        See _reconcile_summary_sheet() for the structure of the result.
        """
        return self._last_summary_reconciliation

    def _validate_sheet_name(self, sheet_name: str) -> bool:
        """Validate that the sheet name is acceptable for processing.
        
        Returns:
            True if the sheet should be processed, False if it should be skipped.
        """
        normalized = sheet_name.strip().lower()

        if normalized == "summary":
            return False

        if not sheet_name.strip():
            raise ValueError("Empty sheet name encountered")

        return True

    def _is_empty_row(self, row: tuple[Any, ...]) -> bool:
        """Check if a row is empty or contains only None values"""
        return all(cell is None or str(cell).strip() == "" for cell in row)

    def _find_header_row(self, sheet_rows: list[tuple[Any, ...]]) -> int:
        """Find the actual student-table header row in a section sheet.

        The source Masterlist workbook places a title row and a blank row
        before the column headers. Do not assume the first worksheet row is
        the header; locate it by the stable field names used by the masterlist.
        """
        required_headers = {"student number", "last name", "first name"}

        for row_index, row in enumerate(sheet_rows):
            normalized_headers = {
                str(cell).strip().lower()
                for cell in row
                if cell is not None and str(cell).strip()
            }
            if required_headers.issubset(normalized_headers):
                return row_index

        raise ValueError(
            "Could not locate the student-table header row in section sheet. "
            "Expected headers including 'Student Number', 'Last Name', and "
            "'First Name'."
        )

    def _extract_row_data(
        self, sheet_name: str, row_idx: int, row: tuple[Any, ...], header_row: tuple[Any, ...]
    ) -> dict[str, Any]:
        """Extract and normalize data from a single row"""
        data = {
            "source_sheet": sheet_name,
            "source_row": row_idx,
            "source_no": None,
            "raw_student_number": None,
            "raw_last_name": None,
            "raw_first_name": None,
            "raw_middle_name": None,
            "raw_section": None,
            "raw_status": None,
            "raw_program": None,
            "raw_year_level": None,
            "raw_academic_year": None,
            "raw_semester": None,
            "raw_subjects_enrolled": None,
            "raw_mobile_number": None,
            "raw_email": None,
        }

        if not header_row:
            return data

        known_headers = [
            "no",
            "student number",
            "last name",
            "first name",
            "middle name",
            "section",
            "status",
            "program",
            "year level",
            "academic year",
            "semester",
            "subjects enrolled",
            "mobile number",
            "email",
        ]

        header_to_col_idx: dict[str, int] = {}
        for col_idx, header_value in enumerate(header_row):
            if header_value is not None:
                header_name = str(header_value).strip().lower()
                for known_header in known_headers:
                    if known_header not in header_to_col_idx and (
                        header_name == known_header or header_name == f"{known_header}s"
                    ):
                        header_to_col_idx[known_header] = col_idx

        col_lookup = {
            "source_no": "no",
            "raw_student_number": "student number",
            "raw_last_name": "last name",
            "raw_first_name": "first name",
            "raw_middle_name": "middle name",
            "raw_section": "section",
            "raw_status": "status",
            "raw_program": "program",
            "raw_year_level": "year level",
            "raw_academic_year": "academic year",
            "raw_semester": "semester",
            "raw_subjects_enrolled": "subjects enrolled",
            "raw_mobile_number": "mobile number",
            "raw_email": "email",
        }

        for field_name, col_key in col_lookup.items():
            col_idx = header_to_col_idx.get(col_key)
            if col_idx is not None and col_idx < len(row):
                data[field_name] = self._safe_cell_value(row[col_idx])

        return data

    def _safe_cell_value(self, cell_value: Any) -> str | None:
        """Safely extract string value from Excel cell, handling None and formatting"""
        if cell_value is None:
            return None

        if isinstance(cell_value, (int, float)):
            return str(int(cell_value) if isinstance(cell_value, int) else cell_value)

        if isinstance(cell_value, str):
            return cell_value.strip()

        return str(cell_value)

    def _derive_program_from_sheet(self, sheet_name: str) -> str:
        """Derive the academic program identity encoded in a section sheet name.

        Per docs/Phase-3-Student-Data-Field-Mapping.md section 4, the workbook
        encodes program through its section sheets using a leading program
        code followed by a year-level/section suffix, e.g. "BSCS 1A" or
        "BSIT 3A". Only the leading program code identifies the program;
        the remainder of the sheet name (year level, section letter)
        identifies the section, not the program, and must not be used to
        distinguish program identity.

        This is a lightweight, sheet-name-only derivation used solely to
        classify import conflicts. It intentionally does not read or rely on
        the parsed `raw_program` field, and does not touch `AcademicProgram`/
        `AcademicSection` records.
        """
        match = re.match(r"^\s*([A-Za-z]+)", sheet_name)
        if match:
            return match.group(1).upper()
        return sheet_name.strip().upper()

    def _derive_section_from_sheet(self, sheet_name: str) -> tuple[str, str, str]:
        """Parse a section sheet name into its program code, year level, and
        section name.

        Sheet names take the generic form
        "<PROGRAM_CODE> <YEAR_LEVEL><SECTION_LETTERS>", where PROGRAM_CODE is
        an alphabetic program code (e.g. "BSCS", "BSIT", or any other
        AcademicProgram.program_code such as a synthetic "PROA"), YEAR_LEVEL
        is one or more digits, and SECTION_LETTERS is one or more letters:

            "BSCS 1A"       -> ("BSCS", "1", "A")
            "BSIT 3A"       -> ("BSIT", "3", "A")
            "  BSCS   2B  " -> ("BSCS", "2", "B")
            "PROA 1A"       -> ("PROA", "1", "A")

        The program code is returned unchanged (after normalization) so it
        can be matched directly against AcademicProgram.program_code.

        Raises ValueError if the sheet name does not match this shape (e.g.
        the "Summary" sheet, or a name missing the program code, year level,
        or section letters).
        """
        normalized = " ".join(sheet_name.strip().split()).upper()
        match = re.fullmatch(r"([A-Z]+)\s+(\d+)([A-Z]+)", normalized)
        if not match:
            raise ValueError(
                f"Cannot derive program/year/section from sheet name "
                f"{sheet_name!r}. Expected a sheet name of the form "
                f"'<PROGRAM_CODE> <YEAR_LEVEL><SECTION_LETTERS>', e.g. "
                f"'BSCS 1A'."
            )

        program_code, year_level, section_name = match.groups()
        return program_code, year_level, section_name

    def _normalize_status(self, raw_status: str | None) -> str | None:
        """Interpret a section sheet's supplied Status value.

        Returns "regular", "irregular", or None when the supplied value is
        missing, blank, or does not clearly represent either. Per
        docs/Phase-3-Student-Data-Field-Mapping.md section 3, this only
        reads the value the workbook actually supplied; it never infers or
        recomputes academic status from subjects or other fields.
        """
        if not raw_status:
            return None

        normalized = raw_status.strip().lower()

        if normalized == "regular":
            return "regular"
        if normalized == "irregular":
            return "irregular"

        return None

    def _safe_numeric(self, cell_value: Any) -> int | None:
        """Best-effort extraction of a whole-number count from a Summary
        sheet cell. Returns None if the cell does not hold a parseable
        whole number (booleans and non-integer floats are deliberately
        rejected rather than silently rounded)."""
        if cell_value is None or isinstance(cell_value, bool):
            return None

        if isinstance(cell_value, int):
            return cell_value

        if isinstance(cell_value, float):
            return int(cell_value) if cell_value.is_integer() else None

        if isinstance(cell_value, str):
            stripped = cell_value.strip()
            if re.fullmatch(r"-?\d+", stripped):
                return int(stripped)

        return None

    def _classify_summary_label(self, normalized_label: str) -> str | None:
        """Classify a Summary sheet row label into one of the four
        reconciliation metrics. Checked in this order so that "Irregular"
        is never mistakenly classified as "Regular" (it contains that
        substring), and "Sections" is never mistakenly classified as a
        total (a "Total Sections" label should reconcile against the
        section count, not the student total)."""
        if "irregular" in normalized_label:
            return "irregular"
        if "regular" in normalized_label:
            return "regular"
        if "section" in normalized_label:
            return "sections"
        if "total" in normalized_label:
            return "total_students"
        return None

    def _read_summary_sheet(
        self, summary_rows: list[tuple[Any, ...]]
    ) -> dict[str, int | None]:
        """Dynamically extract the Summary sheet's declared reconciliation
        figures. Never hardcodes the known example totals - every value is
        read from whatever label/value pairs are actually present on the
        sheet. Supports a label-in-one-cell, value-in-a-later-cell-of-the-
        same-row layout, which is the Summary sheet's stable structure."""
        declared: dict[str, int | None] = {
            "total_students": None,
            "regular": None,
            "irregular": None,
            "sections": None,
        }

        for row in summary_rows:
            if not row:
                continue

            label_idx: int | None = None
            normalized_label: str | None = None
            for idx, cell in enumerate(row):
                if isinstance(cell, str) and cell.strip():
                    label_idx = idx
                    normalized_label = cell.strip().lower().rstrip(":").strip()
                    break

            if normalized_label is None or label_idx is None:
                continue

            metric = self._classify_summary_label(normalized_label)
            if metric is None or declared[metric] is not None:
                continue

            value: int | None = None
            for cell in row[label_idx + 1 :]:
                value = self._safe_numeric(cell)
                if value is not None:
                    break

            if value is not None:
                declared[metric] = value

        return declared

    def _calculate_section_totals(
        self,
        parsed_rows: list[dict[str, Any]],
        processed_section_sheets: list[str],
    ) -> dict[str, int]:
        """Calculate the reconciliation figures from the section-sheet data
        actually discovered by the importer (never from the Summary sheet
        itself)."""
        regular = 0
        irregular = 0
        unknown_status = 0

        for row_data in parsed_rows:
            status = self._normalize_status(row_data.get("raw_status"))
            if status == "regular":
                regular += 1
            elif status == "irregular":
                irregular += 1
            else:
                unknown_status += 1

        return {
            "total_students": len(parsed_rows),
            "regular": regular,
            "irregular": irregular,
            "unknown_status": unknown_status,
            "sections": len(processed_section_sheets),
        }

    def _reconcile_summary_sheet(
        self,
        summary_sheet_rows: list[tuple[Any, ...]] | None,
        parsed_rows: list[dict[str, Any]],
        processed_section_sheets: list[str],
    ) -> dict[str, Any]:
        """3C-04 Summary Reconciliation: compare the Summary sheet's
        declared totals against the section-sheet data discovered during
        import, and report any discrepancies explicitly rather than
        silently proceeding.

        Returns a dict shaped like:
            {
                "summary_sheet_found": bool,
                "status": "matched" | "mismatched" | "unavailable",
                "checks": {
                    "total_students": {
                        "declared": int | None,
                        "calculated": int,
                        "status": "matched" | "mismatched" | "unavailable",
                        "detail": str | None,
                    },
                    "regular": {...},
                    "irregular": {...},
                    "sections": {...},
                },
                "discrepancies": [str, ...],
            }

        "status" is "matched" only when the Summary sheet was found and
        every individual check matched. Any single mismatch makes the
        overall status "mismatched"; if nothing mismatched but a figure
        could not be reconciled (Summary sheet missing, a figure missing
        from it, or section-sheet Status values that are neither Regular
        nor Irregular), the overall status is "unavailable".
        """
        if summary_sheet_rows is None:
            return {
                "summary_sheet_found": False,
                "status": "unavailable",
                "checks": {},
                "discrepancies": ["Summary sheet not found in workbook."],
            }

        declared = self._read_summary_sheet(summary_sheet_rows)
        calculated = self._calculate_section_totals(parsed_rows, processed_section_sheets)

        checks: dict[str, dict[str, Any]] = {}
        discrepancies: list[str] = []

        def add_check(
            metric: str,
            label: str,
            declared_value: int | None,
            calculated_value: int,
            unavailable_reason: str | None = None,
        ) -> None:
            if unavailable_reason is None and declared_value is None:
                unavailable_reason = (
                    f"Could not locate a declared {label.lower()} figure on "
                    f"the Summary sheet."
                )

            if unavailable_reason is not None:
                checks[metric] = {
                    "declared": declared_value,
                    "calculated": calculated_value,
                    "status": "unavailable",
                    "detail": unavailable_reason,
                }
                discrepancies.append(f"{label}: {unavailable_reason}")
                return

            if declared_value == calculated_value:
                checks[metric] = {
                    "declared": declared_value,
                    "calculated": calculated_value,
                    "status": "matched",
                    "detail": None,
                }
                return

            detail = (
                f"Summary declares {declared_value} but {calculated_value} "
                f"were found across the section sheets."
            )
            checks[metric] = {
                "declared": declared_value,
                "calculated": calculated_value,
                "status": "mismatched",
                "detail": detail,
            }
            discrepancies.append(f"{label}: {detail}")

        add_check(
            "total_students",
            "Total students",
            declared["total_students"],
            calculated["total_students"],
        )

        unknown_status_reason: str | None = None
        if calculated["unknown_status"] > 0:
            unknown_status_reason = (
                f"{calculated['unknown_status']} section-sheet row(s) have a "
                f"Status value that is neither Regular nor Irregular, so "
                f"this figure cannot be reliably reconciled."
            )

        add_check(
            "regular",
            "Regular students",
            declared["regular"],
            calculated["regular"],
            unavailable_reason=unknown_status_reason,
        )
        add_check(
            "irregular",
            "Irregular students",
            declared["irregular"],
            calculated["irregular"],
            unavailable_reason=unknown_status_reason,
        )
        add_check(
            "sections",
            "Sections",
            declared["sections"],
            calculated["sections"],
        )

        if any(check["status"] == "mismatched" for check in checks.values()):
            overall_status = "mismatched"
        elif any(check["status"] == "unavailable" for check in checks.values()):
            overall_status = "unavailable"
        else:
            overall_status = "matched"

        return {
            "summary_sheet_found": True,
            "status": overall_status,
            "checks": checks,
            "discrepancies": discrepancies,
        }

    def _validate_and_detect_conflicts(self, parsed_rows: list[dict[str, Any]]) -> None:
        """Validate rows and detect duplicate/cross-program conflicts across the entire workbook"""
        student_to_sheets: dict[str, dict[str, int]] = {}
        student_to_rows: dict[str, dict[str, dict[str, Any]]] = {}
        student_to_all_rows: dict[str, list[dict[str, Any]]] = {}

        for row_data in parsed_rows:
            student_number = row_data.get("raw_student_number")
            if not student_number:
                row_data["validation_status"] = ImportValidationStatus.invalid
                row_data["validation_errors"] = ["Missing student number"]
                row_data["conflict_key"] = None
                continue

            student_number = student_number.strip()

            if student_number not in student_to_sheets:
                student_to_sheets[student_number] = {}
                student_to_rows[student_number] = {}

            sheet_name = row_data["source_sheet"]
            if sheet_name not in student_to_sheets[student_number]:
                student_to_sheets[student_number][sheet_name] = 0
            student_to_sheets[student_number][sheet_name] += 1

            if sheet_name not in student_to_rows[student_number]:
                student_to_rows[student_number][sheet_name] = {}
            student_to_rows[student_number][sheet_name][row_data["source_row"]] = row_data

            if student_number not in student_to_all_rows:
                student_to_all_rows[student_number] = []
            student_to_all_rows[student_number].append(row_data)

        for row_data in parsed_rows:
            student_number = row_data.get("raw_student_number")
            if not student_number:
                continue

            student_number = student_number.strip()

            if student_number not in student_to_sheets:
                continue

            first_name = row_data.get("raw_first_name")
            last_name = row_data.get("raw_last_name")

            if not first_name or not last_name:
                row_data["validation_status"] = ImportValidationStatus.invalid
                row_data["validation_errors"] = ["Missing required name fields"]
                row_data["conflict_key"] = student_number
                continue

            sheet_counts = student_to_sheets[student_number]
            programs_involved = {
                self._derive_program_from_sheet(sheet) for sheet in sheet_counts.keys()
            }

            if len(programs_involved) > 1:
                # The student number appears under section sheets that belong
                # to more than one program (e.g. BSCS and BSIT) - a genuine
                # cross-program conflict per docs/Phase-3-Student-Data-Field-Mapping.md.
                row_data["validation_status"] = ImportValidationStatus.conflict_cross_program
                row_data["validation_errors"] = [
                    f"Conflict across {len(programs_involved)} programs: "
                    f"{', '.join(sorted(programs_involved))}"
                ]
                row_data["conflict_key"] = student_number
            elif len(sheet_counts) == 1:
                # Same program, same sheet: the first occurrence (in parse
                # order) stays valid, every later occurrence is flagged as a
                # duplicate - mirrors the multi-sheet duplicate handling
                # below (canonical_row = first row encountered for this
                # student), rather than unconditionally overwriting
                # validation_status on every outer-loop pass, which
                # previously let whichever row was processed *last* end up
                # valid instead of the first.
                rows_in_sheet = student_to_rows[student_number].get(list(sheet_counts.keys())[0], {})

                if len(rows_in_sheet) > 1:
                    canonical_row = student_to_all_rows[student_number][0]

                    if row_data is canonical_row:
                        row_data["validation_status"] = ImportValidationStatus.valid
                        row_data["validation_errors"] = []
                        row_data["conflict_key"] = None
                    else:
                        row_data["validation_status"] = ImportValidationStatus.invalid
                        row_data["validation_errors"] = [f"Duplicate student number in import: {student_number}"]
                        row_data["conflict_key"] = student_number
                else:
                    row_data["validation_status"] = ImportValidationStatus.valid
                    row_data["validation_errors"] = []
                    row_data["conflict_key"] = None
            elif len(sheet_counts) > 1:
                # Same program, multiple sheets (e.g. BSCS 1A + BSCS 2A): not a
                # cross-program conflict. Reuse the same duplicate/invalid
                # handling used for same-sheet duplicates above - the first
                # occurrence (in parse order) is kept valid, the rest are
                # flagged as duplicates.
                all_rows_for_student = student_to_all_rows[student_number]
                canonical_row = all_rows_for_student[0]

                if row_data is canonical_row:
                    row_data["validation_status"] = ImportValidationStatus.valid
                    row_data["validation_errors"] = []
                    row_data["conflict_key"] = None
                else:
                    row_data["validation_status"] = ImportValidationStatus.invalid
                    row_data["validation_errors"] = [f"Duplicate student number in import: {student_number}"]
                    row_data["conflict_key"] = student_number
            else:
                row_data["validation_status"] = ImportValidationStatus.valid
                row_data["validation_errors"] = []
                row_data["conflict_key"] = None

    def create_staging_records(
        self, import_batch: ImportBatch, parsed_rows: list[dict[str, Any]]
    ) -> list[StudentImportRecord]:
        """Create StudentImportRecord staging records from parsed rows"""
        records = []

        for row_data in parsed_rows:
            record = StudentImportRecord(
                import_batch_id=import_batch.id,
                source_sheet=row_data["source_sheet"],
                source_row=row_data["source_row"],
                source_no=row_data.get("source_no"),
                raw_student_number=row_data.get("raw_student_number"),
                raw_last_name=row_data.get("raw_last_name"),
                raw_first_name=row_data.get("raw_first_name"),
                raw_middle_name=row_data.get("raw_middle_name"),
                raw_mobile_number=row_data.get("raw_mobile_number"),
                raw_email=row_data.get("raw_email"),
                raw_subjects_enrolled=row_data.get("raw_subjects_enrolled"),
                raw_status=row_data.get("raw_status"),
                academic_status=None,
                validation_status=row_data["validation_status"],
                validation_errors=row_data["validation_errors"],
                conflict_key=row_data.get("conflict_key"),
            )

            self.session.add(record)
            records.append(record)

        return records

    def get_validation_summary(self, parsed_rows: list[dict[str, Any]]) -> dict[str, Any]:
        """Get validation summary from parsed rows"""
        summary = {
            "total_rows": len(parsed_rows),
            "valid_rows": 0,
            "invalid_rows": 0,
            "conflict_rows": 0,
            "rows_by_status": {},
        }

        for row_data in parsed_rows:
            status = row_data["validation_status"]

            if status == ImportValidationStatus.valid:
                summary["valid_rows"] += 1
            elif status == ImportValidationStatus.invalid:
                summary["invalid_rows"] += 1
            elif status == ImportValidationStatus.conflict_cross_program:
                summary["conflict_rows"] += 1

            summary["rows_by_status"][status] = (
                summary["rows_by_status"].get(status, 0) + 1
            )

        return summary


# Module-level functions for backward compatibility
def parse_student_import(session: Session, import_batch: ImportBatch, xlsx_file: bytes) -> list[dict[str, Any]]:
    """Parse XLSX workbook and create StudentImportRecord staging rows.

    Returns:
        List of parsed row data with validation results
    """
    service = StudentImportService(session)
    return service.parse_student_import(import_batch, xlsx_file)


def validate_student_import(session: Session, import_batch: ImportBatch, xlsx_file: bytes) -> dict[str, Any]:
    """Parse and validate XLSX workbook, returning validation summary."""
    service = StudentImportService(session)
    parsed_rows = service.parse_student_import(import_batch, xlsx_file)
    return service.get_validation_summary(parsed_rows)


# Export types for type hints
from typing import TypedDict


class ImportRowData(TypedDict):
    source_sheet: str
    source_row: int
