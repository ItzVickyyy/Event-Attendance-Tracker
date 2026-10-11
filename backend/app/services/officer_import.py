"""CSV/XLSX parsing for the two-step Student Council officer import workflow."""

import csv
import io
from typing import Any

from openpyxl import load_workbook  # type: ignore[import-untyped]


ALIASES = {
    "student id": "student_number",
    "student number": "student_number",
    "first name": "first_name",
    "middle initial": "middle_initial",
    "last name": "last_name",
    "extension": "extension",
    "name extension": "extension",
    "academic year": "academic_year",
    "section": "section",
    "email": "email",
    "organization": "organization",
    "role": "role",
    "system role": "role",
    "position": "position",
    "officer position": "position",
}
REQUIRED_COLUMNS = {
    "student_number",
    "first_name",
    "last_name",
    "academic_year",
    "section",
    "email",
    "role",
    "position",
}


def parse_officer_import_file(file_bytes: bytes, filename: str) -> list[dict[str, Any]]:
    extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if extension == "csv":
        try:
            reader = csv.reader(io.StringIO(file_bytes.decode("utf-8-sig")))
            rows = list(reader)
        except UnicodeDecodeError as exc:
            raise ValueError("CSV file must be UTF-8 encoded.") from exc
    elif extension == "xlsx":
        try:
            workbook = load_workbook(
                filename=io.BytesIO(file_bytes), read_only=True, data_only=True
            )
            worksheet = workbook.active
            rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
            workbook.close()
        except Exception as exc:
            raise ValueError(f"Failed to parse XLSX file: {exc}") from exc
    else:
        raise ValueError("Unsupported file type. Upload a .xlsx or .csv file.")

    if not rows:
        raise ValueError("The uploaded file is empty.")
    headers = [_normalize_header(value) for value in rows[0]]
    mapped_headers = [ALIASES.get(header, "") for header in headers]
    available = {header for header in mapped_headers if header}
    missing = sorted(REQUIRED_COLUMNS - available)
    if missing:
        labels = ", ".join(sorted(missing))
        raise ValueError(f"Missing required columns: {labels}")

    parsed: list[dict[str, Any]] = []
    for row_number, cells in enumerate(rows[1:], start=2):
        values = list(cells)
        if all(value is None or str(value).strip() == "" for value in values):
            continue
        record: dict[str, Any] = {"source_row": row_number}
        for index, key in enumerate(mapped_headers):
            if not key or index >= len(values):
                continue
            value = values[index]
            record[key] = "" if value is None else str(value).strip()
        parsed.append(record)
    if not parsed:
        raise ValueError("The uploaded file contains no officer rows.")
    if len(parsed) > 500:
        raise ValueError("Import files are limited to 500 officer accounts at a time.")
    return parsed


def _normalize_header(value: Any) -> str:
    return " ".join(str(value or "").strip().lower().replace("_", " ").split())
